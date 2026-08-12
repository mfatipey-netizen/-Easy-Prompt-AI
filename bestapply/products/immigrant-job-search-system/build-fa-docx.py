#!/usr/bin/env python3
"""Build a print-ready RTL DOCX from a Persian chapter .fa.md.

Usage:
    python3 build-fa-docx.py chapters/01-the-250-to-1-wall.fa.md

Produces:
    chapters/01-the-250-to-1-wall.fa.docx

Pandoc handles Markdown → DOCX conversion (paragraphs, headings, tables,
footnotes). This script then post-processes the DOCX so every paragraph
in the body AND in the footnotes has:

  - <w:bidi/>            → paragraph direction is right-to-left
  - <w:jc w:val="right"/> → paragraph text is right-aligned
  - <w:rtl/> on every run → runs inside the paragraph render RTL

Word / LibreOffice / Google Docs need BOTH bidi and jc="right" to display
Persian text the way the user expects when they open the file to prepare
for print. Pandoc alone sets one of them inconsistently; the two-pass
approach here is what makes the file print-ready.
"""

import io
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{" + W_NS + "}"


def pandoc_convert(md_path: Path, docx_path: Path) -> None:
    """Run pandoc with Persian metadata."""
    subprocess.run(
        [
            "pandoc",
            str(md_path),
            "-o",
            str(docx_path),
            "--metadata=lang:fa-IR",
            "--metadata=dir:rtl",
            "--standalone",
        ],
        check=True,
    )


def rtl_paragraph(p) -> None:
    """Force a python-docx paragraph to RTL and right-align."""
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(OxmlElement("w:bidi"))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        rPr = run._r.get_or_add_rPr()
        if rPr.find(qn("w:rtl")) is None:
            rPr.append(OxmlElement("w:rtl"))


def rtl_table(t) -> None:
    """Force a python-docx table to bidi-visual + right-align every cell."""
    tblPr = t._tbl.find(qn("w:tblPr"))
    if tblPr is not None and tblPr.find(qn("w:bidiVisual")) is None:
        tblPr.append(OxmlElement("w:bidiVisual"))
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                rtl_paragraph(p)


def rtl_body(docx_path: Path) -> None:
    """First pass: body via python-docx."""
    doc = Document(docx_path)
    for p in doc.paragraphs:
        rtl_paragraph(p)
    for t in doc.tables:
        rtl_table(t)
    doc.save(docx_path)


def rtl_footnotes(docx_path: Path) -> None:
    """Second pass: footnotes.xml direct XML manipulation.

    python-docx doesn't expose footnotes cleanly, and the Package
    part-relationship helper varies by python-docx version, so we edit
    the XML in the zip directly. This is boring and reliable.
    """
    ET.register_namespace("w", W_NS)
    ET.register_namespace(
        "r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    with zipfile.ZipFile(docx_path) as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}

    if "word/footnotes.xml" not in parts:
        return

    root = ET.fromstring(parts["word/footnotes.xml"])
    for p in root.iter(W + "p"):
        pPr = p.find(W + "pPr")
        if pPr is None:
            pPr = ET.Element(W + "pPr")
            p.insert(0, pPr)
        if pPr.find(W + "bidi") is None:
            ET.SubElement(pPr, W + "bidi")
        jc = pPr.find(W + "jc")
        if jc is None:
            jc = ET.SubElement(pPr, W + "jc")
        jc.set(W + "val", "right")
        for r in p.iter(W + "r"):
            rPr = r.find(W + "rPr")
            if rPr is None:
                rPr = ET.Element(W + "rPr")
                r.insert(0, rPr)
            if rPr.find(W + "rtl") is None:
                ET.SubElement(rPr, W + "rtl")

    buf = io.BytesIO()
    ET.ElementTree(root).write(buf, xml_declaration=True, encoding="UTF-8")
    parts["word/footnotes.xml"] = buf.getvalue()

    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    shutil.move(tmp, docx_path)


def main() -> None:
    if len(sys.argv) != 2:
        print("usage: build-fa-docx.py path/to/chapter.fa.md", file=sys.stderr)
        sys.exit(2)
    md_path = Path(sys.argv[1]).resolve()
    if not md_path.exists():
        print(f"not found: {md_path}", file=sys.stderr)
        sys.exit(1)
    docx_path = md_path.with_suffix(".docx")
    pandoc_convert(md_path, docx_path)
    rtl_body(docx_path)
    rtl_footnotes(docx_path)
    print(f"built {docx_path}")


if __name__ == "__main__":
    main()
